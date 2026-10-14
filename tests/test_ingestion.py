"""Phase 1 测试：多格式解析、清单差异比对、上传安全校验。

覆盖的边界都是「出过问题或容易出问题」的地方：
  编码回退（GB18030）、PDF 页码标记、DOCX 标题/表格转换、
  CSV 表头语义保留、超大文件、路径穿越。
"""

from __future__ import annotations

import io
import shutil
import tempfile
import unittest
from pathlib import Path

from app.ingestion.loaders import (
    MAX_UPLOAD_BYTES,
    DocumentTooLargeError,
    UnsupportedFormatError,
    detect_file_type,
    page_of,
    parse_bytes,
    strip_page_markers,
)
from app.ingestion.manifest import IndexManifest, ManifestDiff, content_hash


class LoaderTests(unittest.TestCase):
    def test_rejects_unsupported_format(self):
        with self.assertRaises(UnsupportedFormatError):
            detect_file_type("archive.zip")

    def test_rejects_oversized_file(self):
        raw = b"x" * (MAX_UPLOAD_BYTES + 1)
        with self.assertRaises(DocumentTooLargeError):
            parse_bytes(raw, "big.txt")

    def test_markdown_and_text(self):
        md = parse_bytes("# 标题\n\n正文内容".encode("utf-8"), "doc.md")
        self.assertEqual(md.file_type, "markdown")
        self.assertIn("正文内容", md.text)

        txt = parse_bytes("纯文本".encode("utf-8"), "note.txt")
        self.assertEqual(txt.file_type, "text")

    def test_gb18030_fallback(self):
        raw = "中文编码测试".encode("gb18030")
        parsed = parse_bytes(raw, "gbk.txt")
        self.assertIn("中文编码测试", parsed.text)

    def test_csv_becomes_markdown_table(self):
        raw = "波长,温度,点数\n550,25,396\n".encode("utf-8")
        parsed = parse_bytes(raw, "grid.csv")
        self.assertEqual(parsed.file_type, "csv")
        self.assertIn("| 波长 | 温度 | 点数 |", parsed.text)
        self.assertIn("| 550 | 25 | 396 |", parsed.text)

    def test_pdf_page_markers(self):
        # 构造一个最小可用 PDF，验证页码标记与回溯
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        parsed = parse_bytes(buffer.getvalue(), "scan.pdf")
        self.assertEqual(parsed.page_count, 1)
        self.assertIn("[第 1 页]", parsed.text)
        self.assertEqual(page_of("[第 3 页]\n正文"), 3)
        self.assertEqual(strip_page_markers("[第 3 页]\n正文"), "正文")
        self.assertTrue(any("未抽取到文本" in w for w in parsed.warnings))

    def test_docx_heading_becomes_markdown(self):
        import docx

        document = docx.Document()
        document.add_heading("实验方法", level=1)
        document.add_paragraph("这是一段说明文字。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "参数"
        table.cell(0, 1).text = "数值"
        table.cell(1, 0).text = "噪声"
        table.cell(1, 1).text = "0.01"
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = parse_bytes(buffer.getvalue(), "report.docx")
        self.assertIn("# 实验方法", parsed.text)
        self.assertIn("这是一段说明文字。", parsed.text)
        self.assertIn("| 参数 | 数值 |", parsed.text)


class ManifestTests(unittest.TestCase):
    def test_content_hash_is_stable_and_short(self):
        digest = content_hash(b"hello")
        self.assertEqual(digest, content_hash(b"hello"))
        self.assertEqual(len(digest), 16)
        self.assertNotEqual(digest, content_hash(b"hello!"))

    def test_diff_detects_added_changed_removed(self):
        manifest = IndexManifest(
            embed_model="m",
            chunk_config={"a": 1},
            files=[
                {"rel_path": "a.md", "content_hash": "1", "size": 1, "file_type": "markdown"},
                {"rel_path": "b.md", "content_hash": "2", "size": 1, "file_type": "markdown"},
            ],
        )
        current = {
            "a.md": {"rel_path": "a.md", "content_hash": "1", "size": 1, "file_type": "markdown"},
            "b.md": {"rel_path": "b.md", "content_hash": "9", "size": 2, "file_type": "markdown"},
            "c.md": {"rel_path": "c.md", "content_hash": "3", "size": 1, "file_type": "markdown"},
        }
        diff = manifest.diff_against(current)
        self.assertEqual(diff.added, ["c.md"])
        self.assertEqual(diff.changed, ["b.md"])
        self.assertEqual(diff.removed, [])
        self.assertEqual(diff.unchanged, ["a.md"])
        self.assertTrue(diff.needs_rebuild)

    def test_compatibility_rejects_changed_model_or_chunking(self):
        manifest = IndexManifest(embed_model="bge-small", chunk_config={"child_size": 500})
        ok, _ = manifest.is_compatible("bge-small", {"child_size": 500})
        self.assertTrue(ok)
        ok, reason = manifest.is_compatible("bge-large", {"child_size": 500})
        self.assertFalse(ok)
        self.assertIn("嵌入模型", reason)
        ok, reason = manifest.is_compatible("bge-small", {"child_size": 800})
        self.assertFalse(ok)
        self.assertIn("切块", reason)

    def test_manifest_roundtrip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "manifest.json"
            manifest = IndexManifest(built_at="2026-01-01T00:00:00", child_count=7)
            manifest.save(path)
            loaded = IndexManifest.load(path)
            self.assertIsNotNone(loaded)
            self.assertEqual(loaded.child_count, 7)
            self.assertEqual(loaded.built_at, "2026-01-01T00:00:00")

    def test_diff_summary_when_clean(self):
        diff = ManifestDiff(added=[], changed=[], removed=[], unchanged=["a.md"])
        self.assertFalse(diff.needs_rebuild)
        self.assertIn("无变化", diff.summary())


class SaveUploadTests(unittest.TestCase):
    def _settings(self, tmp: Path):
        class FakeSettings:
            docs_dir = tmp / "docs"

        return FakeSettings()

    def test_rejects_path_traversal(self):
        from app.ingestion.indexer import save_upload

        with tempfile.TemporaryDirectory() as tmp:
            settings = self._settings(Path(tmp))
            settings.docs_dir.mkdir(parents=True)
            ok, message = save_upload(settings, "../escape.md", b"x")
            self.assertFalse(ok)
            self.assertIn("非法", message)
            self.assertFalse((Path(tmp) / "escape.md").exists())

    def test_accepts_valid_upload_and_creates_subdir(self):
        from app.ingestion.indexer import save_upload

        with tempfile.TemporaryDirectory() as tmp:
            settings = self._settings(Path(tmp))
            settings.docs_dir.mkdir(parents=True)
            ok, _ = save_upload(settings, "论文/new.md", b"# hello")
            self.assertTrue(ok)
            self.assertTrue((settings.docs_dir / "论文" / "new.md").exists())

    def test_rejects_unsupported_and_empty(self):
        from app.ingestion.indexer import save_upload

        with tempfile.TemporaryDirectory() as tmp:
            settings = self._settings(Path(tmp))
            settings.docs_dir.mkdir(parents=True)
            self.assertFalse(save_upload(settings, "x.exe", b"data")[0])
            self.assertFalse(save_upload(settings, "x.md", b"")[0])


if __name__ == "__main__":
    unittest.main()
